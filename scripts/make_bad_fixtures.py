#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""Generate test fixtures with intentional data issues for testing."""

from pathlib import Path
import shutil
from openpyxl import load_workbook

def make_zaiko_ma_rong():
    """Create zaiko_ma_rong.xlsx by blanking some warehouse_code fields."""
    source = Path("tests/fixtures/zaiko_ok.xlsx")
    target = Path("tests/fixtures/zaiko_ma_rong.xlsx")

    # Copy source to target
    shutil.copy(source, target)

    # Load and modify
    wb = load_workbook(target)
    ws = wb.active

    # Blank out warehouse_code for rows 2, 3, 5
    # Header is at row 1, so data starts at row 2
    column_idx = None
    for col_idx, cell in enumerate(ws[1], 1):
        if cell.value and "倉庫コード" in str(cell.value):
            column_idx = col_idx
            break

    if column_idx:
        # Blank out a few rows
        for row_idx in [2, 3, 5]:
            ws.cell(row=row_idx, column=column_idx).value = None

    wb.save(target)
    print("Created zaiko_ma_rong.xlsx with empty warehouse codes")

def make_zaiko_thieu_cot():
    """Create zaiko_thieu_cot.xlsx by deleting the stock_qty column."""
    source = Path("tests/fixtures/zaiko_ok.xlsx")
    target = Path("tests/fixtures/zaiko_thieu_cot.xlsx")

    # Load source
    wb = load_workbook(source)
    ws = wb.active

    # Find and delete the stock_qty column
    column_idx = None
    for col_idx, cell in enumerate(ws[1], 1):
        if cell.value and "在庫残数" in str(cell.value):
            column_idx = col_idx
            break

    if column_idx:
        ws.delete_cols(column_idx)

    wb.save(target)
    print("Created zaiko_thieu_cot.xlsx by deleting stock_qty column")

def make_zaiko_cat_cut():
    """Create zaiko_cat_cut.xlsx with only ~10 rows of data."""
    source = Path("tests/fixtures/zaiko_ok.xlsx")
    target = Path("tests/fixtures/zaiko_cat_cut.xlsx")

    # Load source
    wb = load_workbook(source)
    ws = wb.active

    # Delete rows 12 onwards (keep header + 10 data rows)
    if ws.max_row > 11:
        ws.delete_rows(12, ws.max_row)

    wb.save(target)
    print("Created zaiko_cat_cut.xlsx with only 10 data rows")

if __name__ == "__main__":
    make_zaiko_ma_rong()
    make_zaiko_thieu_cot()
    make_zaiko_cat_cut()

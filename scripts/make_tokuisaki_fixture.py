"""Sinh fixture 得意先全情報 ĐÃ ẨN DANH từ bản xuất thật của OBC.

Giữ nguyên: toàn bộ 317 cột, mọi cột mã, mã bưu chính, tỉnh/thành, hạng, ngành,
người phụ trách. Mã bưu chính GIỮ LẠI vì nó là bẫy số 0 đầu cần được test
(8 khách 北海道 đã bị OBC làm mất số 0 ngay trong bản xuất gốc).
Xoá sạch: tên khách, địa chỉ chi tiết, điện thoại, fax, email, web, mã số hoá đơn, ghi chú.

Lý do: dữ liệu thật chứa 個人情報 (tên, địa chỉ, liên hệ của 2.080 khách),
không được đưa vào git dù repo là private. Fixture chỉ cần cấu trúc và các mã.

Chạy lại khi OBC đổi cấu trúc cột:
    python scripts/make_tokuisaki_fixture.py "<đường dẫn file thật>"
"""
import sys
from pathlib import Path
import openpyxl

SCRUB_EXACT = {
    "得意先名", "得意先名カナ", "事業所名", "事業所名カナ", "得意先略称",
    "インデックス", "番地", "ビル等", "電話番号", "ＦＡＸ番号",
    "ホームページ", "メモ２", "メモ３", "インボイス登録番号", "法人番号",
    "取引先名", "取引先事業所名", "請求先名",
}
SCRUB_CONTAINS = ("ご担当", "担当者名（", "携帯", "Ｅ－Ｍａｉｌ", "E-Mail", "メール")

def main(src: Path, dst: Path) -> None:
    wb = openpyxl.load_workbook(src)
    ws = wb.worksheets[0]
    header = [c.value for c in ws[1]]
    scrub = [
        i for i, h in enumerate(header)
        if h and (str(h) in SCRUB_EXACT or any(k in str(h) for k in SCRUB_CONTAINS))
    ]
    for r, row in enumerate(ws.iter_rows(min_row=2), start=2):
        for i in scrub:
            row[i].value = None
        # Đặt tên giả ổn định để vẫn phân biệt được các dòng
        row[header.index("得意先名")].value = f"KHACH_{r - 1:05d}"
    wb.save(dst)
    print(f"đã ẩn danh {len(scrub)} cột, ghi {dst}")

if __name__ == "__main__":
    real = Path(sys.argv[1])
    out = Path(__file__).resolve().parent.parent / "tests" / "fixtures" / "tokuisaki_ok.xlsx"
    main(real, out)
